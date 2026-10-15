#!/usr/bin/env python3
"""Contrôle des exports, galerie GitHub, aperçu et archive de transmission V7."""
from pathlib import Path
import hashlib, json, runpy, zipfile
import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
ROOT = Path(__file__).resolve().parents[2]
V7 = ROOT/'v7'

def main():
    data=runpy.run_path(str(V7/'scripts/generer_dossier_concours.py'))
    pdf=fitz.open(V7/'DOSSIER_CONCOURS_V7.pdf')
    ppt=Presentation(V7/'DOSSIER_CONCOURS_V7_EDITABLE.pptx')
    assert len(pdf)==len(ppt.slides)==len(data['pages'])
    assert data['IFC_READY'], 'Dessins ou caméras provenant d’un autre IFC'
    assert all('VISUEL EN COURS' not in p.get_text() for p in pdf)
    sha=hashlib.sha256((V7/'DECIMO_PROPOSITION_V7.ifc').read_bytes()).hexdigest()
    assert json.loads((V7/'visite_3d_provenance.json').read_text())['source_sha256']==sha
    for p in (V7/'assets/ifc').glob('*_camera.json'):
        assert json.loads(p.read_text())['source_sha256']==sha
    gallery=['# Bibliothèque Nandjui D. Simon — une bibliothèque habitée',
             f'Les {len(pdf)} planches du dossier V7 sont affichées ci-dessous. Cliquer sur une image pour l’agrandir.',
             '[Télécharger le PDF](DOSSIER_CONCOURS_V7.pdf) · [PowerPoint modifiable](DOSSIER_CONCOURS_V7_EDITABLE.pptx) · [IFC](DECIMO_PROPOSITION_V7.ifc) · [Pack complet](PACK_PROJET_V7.zip)',
             'Les ambiances IA, le photomontage sur site et les dessins extraits de l’IFC sont légendés. Les dimensions de maquette restent à confronter au relevé avant exécution.']
    rows=(len(pdf)+3)//4
    sheet=Image.new('RGB',(1800,rows*347),'#e8e1d8');draw=ImageDraw.Draw(sheet)
    for i,(p,meta) in enumerate(zip(pdf,data['pages'])):
        gallery.extend([f"## {i+1:02d} — {meta['title']}",f"![Planche {i+1:02d} — {meta['title']}](apercus_concours/page_{i+1:02d}.png)"])
        pix=p.get_pixmap(matrix=fitz.Matrix(.85,.85),alpha=False)
        im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
        im.thumbnail((434,308))
        x=(i%4)*450+(450-im.width)//2;y=(i//4)*347+10
        sheet.paste(im,(x,y));draw.text(((i%4)*450+10,(i//4)*347+324),f"{i+1:02d}  {meta['title']}",fill='#292826')
    (V7/'VOIR_LE_DOSSIER.md').write_text('\n\n'.join(gallery)+'\n',encoding='utf-8')
    sheet.save(V7/'APERCU_DOSSIER_V7.jpg',quality=90)
    native_text=sum(s.has_text_frame for sl in ppt.slides for s in sl.shapes)
    check={'ifc_sha256':sha,'pages_pdf':len(pdf),'format_pdf_mm':[round(v*25.4/72,1) for v in pdf[0].rect[2:]],'slides_pptx':len(ppt.slides),'elements_texte_editables_pptx':native_text,'dessins_et_cameras_meme_ifc':True,'geometrie_ifc':'207 produits représentés, 207 géométries réussies, 0 échec; voir validation_ifc_v7.txt','contenu':'Conception et présentation; mesures de maquette à confirmer sur site.'}
    (V7/'CONTROLE_LIVRAISON.json').write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n')
    archive=V7/'PACK_PROJET_V7.zip'
    files=[p for p in V7.rglob('*') if p.is_file() and p!=archive and '__pycache__' not in p.parts]
    for folder in [ROOT/'photos',ROOT/'charte',ROOT/'textures',ROOT/'v5/scripts',ROOT/'v5/rapports']:
        files.extend(p for p in folder.rglob('*') if p.is_file() and '__pycache__' not in p.parts)
    files.extend([ROOT/'DECIMO_ESQ_V4.ifc',ROOT/'v5/DECIMO_ESQ_V5.ifc',ROOT/'v6/DECIMO_ESQ_V6_OPTION_C.ifc'])
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(set(files)):z.write(p,p.relative_to(ROOT))
    with zipfile.ZipFile(archive) as z:assert z.testzip() is None
    print(json.dumps(check,ensure_ascii=False,indent=2))
    print('Archive',archive.stat().st_size,'octets;',len(set(files)),'fichiers')

if __name__=='__main__':main()
